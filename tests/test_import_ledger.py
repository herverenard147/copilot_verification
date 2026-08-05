"""src/import_ledger.py : parsing de fichiers réels (.xlsx/.csv/.docx) en
lignes d'écriture. Génère de vrais fichiers en mémoire (openpyxl/python-docx),
pas des mocks -- si le parsing casse sur un vrai fichier Excel/Word, ces tests
le détectent."""
import io

import pytest
from openpyxl import Workbook

from src.import_ledger import parse_ledger_file


def _xlsx_bytes(rows):
    """rows[0] = en-tête, rows[1:] = données."""
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _docx_bytes(rows):
    from docx import Document
    doc = Document()
    table = doc.add_table(rows=0, cols=len(rows[0]))
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_xlsx_basique():
    raw = _xlsx_bytes([
        ["Compte", "Libellé", "Débit", "Crédit"],
        ["101", "Capital", "", "100000"],
        ["521", "Banque", "100000", ""],
    ])
    rows, errors = parse_ledger_file("bilan.xlsx", raw)
    assert errors == []
    assert len(rows) == 2
    assert rows[0] == {"account": "101", "label": "Capital", "debit": 0.0, "credit": 100000.0}
    assert rows[1]["account"] == "521"
    assert rows[1]["debit"] == 100000.0


def test_csv_basique():
    csv = "Compte,Libelle,Debit,Credit\n601,Achats,1000,\n571,Caisse,,1000\n"
    rows, errors = parse_ledger_file("journal.csv", csv.encode("utf-8"))
    assert errors == []
    assert len(rows) == 2
    assert rows[0]["account"] == "601"
    assert rows[0]["debit"] == 1000.0


def test_docx_tableau():
    raw = _docx_bytes([
        ["Compte", "Libellé", "Débit", "Crédit"],
        ["101", "Capital", "", "50 000"],
        ["521", "Banque", "50 000", ""],
    ])
    rows, errors = parse_ledger_file("bilan.docx", raw)
    assert errors == []
    assert len(rows) == 2
    assert rows[0]["credit"] == 50000.0   # espace milliers tolere


def test_virgule_decimale_tolere():
    csv = "Compte,Debit,Credit\n601,\"1 234,56\",\n"
    rows, errors = parse_ledger_file("j.csv", csv.encode("utf-8"))
    assert errors == []
    assert rows[0]["debit"] == 1234.56


def test_colonnes_alias_anglais_acceptees():
    csv = "Account,Debit,Credit\n601,500,\n"
    rows, errors = parse_ledger_file("j.csv", csv.encode("utf-8"))
    assert errors == []
    assert rows[0]["account"] == "601"


def test_ligne_sans_compte_ignoree_mais_le_reste_importe():
    csv = "Compte,Debit,Credit\n,500,\n601,1000,\n"
    rows, errors = parse_ledger_file("j.csv", csv.encode("utf-8"))
    assert len(rows) == 1
    assert rows[0]["account"] == "601"
    assert len(errors) == 1
    assert errors[0]["row"] == 2
    assert "compte" in errors[0]["reason"]


def test_ligne_montant_illisible_ignoree():
    csv = "Compte,Debit,Credit\n601,pas-un-nombre,\n605,500,\n"
    rows, errors = parse_ledger_file("j.csv", csv.encode("utf-8"))
    assert len(rows) == 1
    assert rows[0]["account"] == "605"
    assert len(errors) == 1
    assert "montant" in errors[0]["reason"]


def test_ligne_debit_et_credit_nuls_ignoree():
    csv = "Compte,Debit,Credit\n601,0,0\n605,,\n611,500,\n"
    rows, errors = parse_ledger_file("j.csv", csv.encode("utf-8"))
    assert len(rows) == 1
    assert len(errors) == 2


def test_aucune_colonne_compte_leve_erreur_claire():
    csv = "Foo,Bar\n1,2\n"
    with pytest.raises(ValueError, match="Compte"):
        parse_ledger_file("j.csv", csv.encode("utf-8"))


def test_fichier_vide_leve_erreur():
    csv = "Compte,Debit,Credit\n"
    with pytest.raises(ValueError):
        parse_ledger_file("j.csv", csv.encode("utf-8"))


def test_format_non_reconnu():
    with pytest.raises(ValueError, match="non reconnu"):
        parse_ledger_file("bilan.pdf", b"%PDF-1.4 ...")


def test_fichier_corrompu_ne_plante_pas_avec_trace_python():
    with pytest.raises(ValueError):
        parse_ledger_file("bilan.xlsx", b"ceci n'est pas un vrai xlsx")
